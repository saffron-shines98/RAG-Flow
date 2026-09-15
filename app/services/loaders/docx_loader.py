"""
services/loaders/docx_loader.py
=================================
Word (.docx) files se text nikalna.
"""

from docx import Document


def extract_text(file_path: str) -> str:
    doc = Document(file_path)
    full_text = []

    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                full_text.append(row_text)

    return "\n".join(full_text)